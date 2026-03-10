"""Extract workbook structure, labels, and formulas for mapping/parity updates.

Usage:
  python scripts/extract_excel_mapping.py
  python scripts/extract_excel_mapping.py "docs/source/excel/Coffee Doc-Praxis-V2.xlsm"
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

DEFAULT_WORKBOOK = Path("docs/source/excel/Coffee Doc-Praxis-V2.xlsm")
OUTPUT_JSON = Path("docs/source/excel/extraction-report.json")
OUTPUT_MD = Path("docs/source/excel/extraction-summary.md")

FOCUS_SHEETS = [
    "Form Configuration",
    "Contract",
    "Shipping Instruction",
    "Bank & LC",
    "Contract-SI-LC",
    "Bookings",
    "Staffing",
    "Processing",
    "Commercial Invoice",
    "Packing List",
    "SI",
    "Certificate of Quality",
    "Certificate of Weight",
    "WAY BILL",
    "COCG",
    "VGM",
    "ICO",
    "Cert Of Origin",
    "NON-GMO",
]


def _sheet_profile(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict:
    formulas = []
    labels = []

    max_row = min(ws.max_row, 220)
    max_col = min(ws.max_column, 20)

    for r in range(1, max_row + 1):
        rowvals = []
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            val = cell.value
            if val is None:
                continue
            if isinstance(val, str):
                val = val.strip()
                if not val:
                    continue
            rowvals.append((cell.coordinate, val))
            if isinstance(val, str) and val.startswith("=") and len(formulas) < 200:
                formulas.append({"cell": cell.coordinate, "formula": val})

        if rowvals and len(labels) < 160:
            text_parts = [
                f"{coord}:{str(v).replace(chr(10), ' ')[:120]}"
                for coord, v in rowvals
                if not (isinstance(v, str) and v.startswith("="))
            ]
            if text_parts:
                labels.append(" | ".join(text_parts[:10]))

    formula_count_total = sum(
        1
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )

    return {
        "sheet_state": ws.sheet_state,
        "used_range": {
            "min_row": ws.min_row,
            "max_row": ws.max_row,
            "min_col": ws.min_column,
            "max_col": ws.max_column,
        },
        "formula_count_total_est": formula_count_total,
        "labels_sample": labels,
        "formulas_sample": formulas,
    }


def main() -> None:
    workbook_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    wb = openpyxl.load_workbook(workbook_path, data_only=False, keep_vba=True)

    summary: dict = {
        "workbook": str(workbook_path),
        "sheets": wb.sheetnames,
        "defined_names": {},
        "sheet_profiles": {},
    }

    for key in wb.defined_names.keys():
        defn = wb.defined_names[key]
        try:
            summary["defined_names"][key] = list(defn.destinations)
        except Exception:
            summary["defined_names"][key] = str(getattr(defn, "attr_text", ""))

    for title in FOCUS_SHEETS:
        if title not in wb.sheetnames:
            continue
        summary["sheet_profiles"][title] = _sheet_profile(wb[title])

    OUTPUT_JSON.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")

    lines = [
        "# Workbook Extraction Summary",
        "",
        f"- File: `{workbook_path}`",
        f"- Sheet count: {len(wb.sheetnames)}",
        "",
        "## Sheets",
    ]
    lines.extend([f"- {s}" for s in wb.sheetnames])
    lines.append("")
    lines.append("## Focus Sheet Metrics")

    for title, profile in summary["sheet_profiles"].items():
        rng = profile["used_range"]
        lines.append(f"### {title}")
        lines.append(f"- State: `{profile['sheet_state']}`")
        lines.append(
            f"- Used range: rows {rng['min_row']}-{rng['max_row']}, cols {rng['min_col']}-{rng['max_col']}"
        )
        lines.append(f"- Formula cells: {profile['formula_count_total_est']}")
        lines.append("- Labels sample:")
        lines.extend([f"  - {row}" for row in profile["labels_sample"][:8]])
        if profile["formulas_sample"]:
            lines.append("- Formula sample:")
            lines.extend(
                [f"  - {f['cell']}: `{f['formula']}`" for f in profile["formulas_sample"][:8]]
            )
        lines.append("")

    OUTPUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {OUTPUT_JSON} and {OUTPUT_MD}")


if __name__ == "__main__":
    main()
